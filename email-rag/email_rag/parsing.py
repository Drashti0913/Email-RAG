"""MIME parsing and bounded attachment extraction; never execute email content."""
from dataclasses import dataclass, field
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
import re
import zipfile

MAX_EMAIL_BYTES = 30 * 1024 * 1024
MAX_ATTACHMENT_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.hidden += 1
        if tag in ('p', 'div', 'br', 'li', 'tr'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def html_text(value):
    parser = HTMLText()
    parser.feed(value)
    return '\n'.join(x.strip() for x in ''.join(parser.parts).splitlines() if x.strip())


def decode(payload, charset='utf-8'):
    try:
        return payload.decode(charset or 'utf-8', errors='replace')
    except LookupError:
        return payload.decode('utf-8', errors='replace')


def extract_attachment(payload, filename, content_type, charset='utf-8'):
    """Return readable text or raise; caller records every skipped attachment."""
    if len(payload) > MAX_ATTACHMENT_BYTES:
        raise ValueError('attachment exceeds 12 MiB extraction limit')
    suffix = Path(filename).suffix.lower()
    if content_type.startswith('text/') or suffix in ('.txt', '.md', '.csv', '.json'):
        value = decode(payload, charset)
        result = html_text(value) if suffix in ('.html', '.htm') or content_type == 'text/html' else value
    elif suffix == '.pdf' or content_type == 'application/pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(payload))
        if reader.is_encrypted:
            raise ValueError('encrypted PDF requires decryption before ingest')
        if len(reader.pages) > 200:
            raise ValueError('PDF exceeds 200-page limit')
        result = '\n'.join(p.extract_text() or '' for p in reader.pages)
    elif suffix == '.docx':
        from docx import Document
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('DOCX uncompressed size exceeds 40 MiB')
        doc = Document(BytesIO(payload))
        result = '\n'.join([p.text for p in doc.paragraphs] +
                           [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    else:
        raise ValueError('unsupported format (images/scans need OCR; archives are not extracted)')
    if not result.strip():
        raise ValueError('no extractable text; scanned documents need OCR')
    if len(result) > MAX_TEXT_CHARS:
        raise ValueError('extracted text exceeds 1,000,000 characters')
    return result.strip()


@dataclass
class EmailRecord:
    email_id: str
    content_hash: str
    sender: str
    recipient: str
    subject: str
    sent_at: datetime | None
    parts: list[tuple[str, str]]
    source_url: str = ''
    warnings: list[str] = field(default_factory=list)


def parse_email(raw: bytes, external_id: str | None = None) -> EmailRecord:
    if len(raw) > MAX_EMAIL_BYTES:
        raise ValueError('email exceeds 30 MiB limit')
    message = BytesParser(policy=policy.default).parsebytes(raw)
    identity = external_id or str(message.get('Message-ID') or sha256(raw).hexdigest())
    record = EmailRecord(
        sha256(identity.encode()).hexdigest(), sha256(b'parser-v1:' + raw).hexdigest(),
        str(message.get('From', '')), str(message.get('To', '')),
        str(message.get('Subject', '(no subject)')), None, [])
    record.warnings.extend('MIME defect: ' + type(d).__name__ for d in message.defects)
    try:
        dt = parsedate_to_datetime(str(message.get('Date', '')))
        record.sent_at = (dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)).astimezone(timezone.utc)
    except (ValueError, TypeError, OverflowError):
        record.warnings.append('missing/invalid Date; excluded from latest queries')

    def visit(part):
        filename = part.get_filename()
        attachment = part.get_content_disposition() == 'attachment' or filename is not None
        if attachment:
            name = filename or 'unnamed-attachment'
            try:
                payload = part.get_payload(decode=True) or b''
                text = extract_attachment(payload, name, part.get_content_type(), part.get_content_charset())
                record.parts.append((name, text))
            except Exception as exc:
                record.warnings.append(f'{name}: {type(exc).__name__}: {exc}')
            return
        if part.is_multipart():
            children = list(part.iter_parts())
            if part.get_content_subtype() == 'alternative':
                # Choose plain text over duplicate HTML; preserve nested alternatives.
                selected = next((p for p in children if p.get_content_type() == 'text/plain'), None)
                selected = selected or next((p for p in children if p.get_content_type() == 'text/html'), None)
                if selected is not None:
                    visit(selected)
                elif children:
                    visit(children[-1])
            else:
                for child in children:
                    visit(child)
            return
        if part.get_content_type() in ('text/plain', 'text/html'):
            value = decode(part.get_payload(decode=True) or b'', part.get_content_charset())
            if part.get_content_type() == 'text/html':
                value = html_text(value)
            if value.strip():
                if len(value) > MAX_TEXT_CHARS:
                    raise ValueError('email body exceeds text limit')
                record.parts.append(('body', value.strip()))
    visit(message)
    return record


def chunk_text(text, size=1200, overlap=180):
    """Character chunks retain every character, including long unbroken strings."""
    if not 0 <= overlap < size:
        raise ValueError('require 0 <= overlap < size')
    text = text.replace('\x00', '')
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        yield text[start:end]
        if end == len(text):
            break
        start = end - overlap


def chunks_for(record):
    # Bound headers so adversarial headers cannot monopolize model context.
    header = f'Subject: {record.subject[:300]}\nFrom: {record.sender[:300]}\nTo: {record.recipient[:300]}\nDate: {record.sent_at}\n'
    parts = record.parts or [('body', '(No extractable body; metadata only.)')]
    return [(name, header + f'Part: {name[:200]}\n' + text)
            for name, value in parts for text in chunk_text(value)]
